#!/usr/bin/env python3
# coding: utf-8
"""把技能产出的 Markdown 用例表导出成 xlsx（S6/S8 的交付格式）。

两种列结构：
  generic             给人看/评审：编号 · 接口 · 场景 · 优先级 · 前置 · 步骤 · 预期 · 判据 · 可测性 · 证据/关联
  autotest-platform   给 AutoTestAgent「用例页 → 导入 Excel」直接吃：
                      用例编号 · 用例名称 · 来源 · 前置条件 · 操作步骤 · 预期结果 · 状态 · 脚本状态 · 当前版本 · 更新时间
                      （后接 判据/可测性/优先级/关联接口/验证状态/证据 六个附加列，平台导入时忽略，但人看得到）

实现有两条路，自动选择：
  1. 有 openpyxl → 用它（可设字体、列宽、冻结首行；交付件按 skill 要求用统一专业字体 Arial）
  2. 没有 openpyxl → 走纯标准库（zipfile + XML 手写最小 xlsx），保证任何机器都能出文件

⚠️ 与「平台导入」有关的一个已知事实：平台的导入器只认它自己的那 10 列，
   **`判据` 列会被忽略**（normalizeCaseRow 的别名表里没有"判据"，INSERT 也不写 oracle_json）。
   所以导入后这些用例在平台里仍是"没有可机器校验判据"的状态 —— 要么先给平台导入器加判据列映射，
   要么导入后重新生成判据。导出时脚本会把这句提醒打在 stdout。

用法：
    python export_cases_xlsx.py <用例文件.md> [--out 输出.xlsx] [--profile generic|autotest-platform]

退出码：0 成功；1 参数/解析失败
"""
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent

GENERIC_COLUMNS = ['编号', '接口', '场景', '优先级', '前置', '步骤', '预期', '判据', '可测性', '证据/关联']
PLATFORM_COLUMNS = ['用例编号', '用例名称', '来源', '前置条件', '操作步骤', '预期结果', '状态', '脚本状态', '当前版本', '更新时间']
EXTRA_COLUMNS = ['判据', '可测性', '优先级', '关联接口', '验证状态', '证据']

VERIFIED_RE = re.compile(r'已真机验证通过|已改造并真机验证通过')


# ---------- 解析 Markdown 用例表 ----------

def parse_case_table(md_text):
    """取表头同时含「编号」与「判据」的那张表，返回 (header, rows)。"""
    header, rows = None, []
    for raw in md_text.splitlines():
        line = raw.strip()
        if not (line.startswith('|') and line.endswith('|')):
            if header:
                break
            continue
        cells = [c.strip() for c in line.strip('|').split('|')]
        if all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c):
            continue
        if header is None:
            if '编号' in cells and '判据' in cells:
                header = cells
            continue
        rows.append(cells)
    return header, rows


def cell(row, header, name):
    if name not in header:
        return ''
    i = header.index(name)
    return row[i].strip() if i < len(row) else ''


def steps_of(text):
    return [s.strip() for s in re.split(r'<br\s*/?>', text) if s.strip()]


def build_rows(header, rows, profile):
    """把解析出的用例行转成 (列名列表, 数据行列表)。"""
    if profile == 'generic':
        cols = [c for c in GENERIC_COLUMNS if c in header or c in ('编号', '接口', '场景', '步骤', '判据')]
        data = []
        for r in rows:
            got = {}
            for c in cols:
                v = cell(r, header, c)
                got[c] = '\n'.join(steps_of(v)) if c == '步骤' else v.replace('<br>', '；')
            data.append([got[c] for c in cols])
        return cols, data

    cols = PLATFORM_COLUMNS + EXTRA_COLUMNS
    data = []
    for r in rows:
        case_no = cell(r, header, '编号')
        interface = cell(r, header, '接口')
        scenario = cell(r, header, '场景')
        oracle = cell(r, header, '判据').replace('<br>', '；')
        expected = cell(r, header, '预期')
        evidence = cell(r, header, '证据/关联').replace('<br>', '；')
        steps = '\n'.join(steps_of(cell(r, header, '步骤')))
        data.append([
            case_no,
            f'{interface} · {scenario}场景',
            'AI 生成',
            cell(r, header, '前置'),
            steps,
            f'{expected}（判据：{oracle}）' if oracle else expected,
            '未执行',
            '未绑定',
            1,
            datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            oracle,
            cell(r, header, '可测性'),
            cell(r, header, '优先级'),
            interface,
            '已真机验证通过' if VERIFIED_RE.search(evidence) else '未验证（草案）',
            evidence[:500],
        ])
    return cols, data


# ---------- openpyxl 路径 ----------

def write_with_openpyxl(path, cols, data, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = '测试用例'
    head_font = Font(name='Arial', bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', start_color='305496')
    body_font = Font(name='Arial')
    ws.append(cols)
    for i, _ in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in data:
        ws.append(row)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = body_font
            c.alignment = Alignment(vertical='top', wrap_text=True)
    widths = {'编号': 12, '用例编号': 12, '接口': 24, '用例名称': 34, '关联接口': 24, '场景': 12,
              '优先级': 8, '可测性': 8, '前置': 20, '前置条件': 20, '步骤': 46, '操作步骤': 46,
              '预期': 40, '预期结果': 46, '判据': 42, '证据/关联': 50, '证据': 50,
              '来源': 10, '状态': 10, '脚本状态': 10, '当前版本': 10, '更新时间': 20, '验证状态': 16}
    for i, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 18)
    ws.freeze_panes = 'A2'
    if meta:
        ms = wb.create_sheet('说明')
        ms.append(['项', '值'])
        ms['A1'].font = head_font
        ms['B1'].font = head_font
        ms['A1'].fill = ms['B1'].fill = head_fill
        for k, v in meta:
            ms.append([k, v])
        ms.column_dimensions['A'].width = 22
        ms.column_dimensions['B'].width = 90
        for r in ms.iter_rows(min_row=2, max_row=ms.max_row):
            for c in r:
                c.font = body_font
                c.alignment = Alignment(vertical='top', wrap_text=True)
    wb.save(path)
    return 'openpyxl'


# ---------- 标准库降级路径（手写最小 xlsx） ----------

def _esc(s):
    return (str(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            .replace('"', '&quot;').replace("'", '&apos;'))


def _col_letter(n):
    s = ''
    while n > 0:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


def _sheet_xml(rows):
    out = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
           '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">',
           '<cols>']
    for i, _ in enumerate(rows[0] if rows else [], start=1):
        out.append(f'<col min="{i}" max="{i}" width="22" customWidth="1"/>')
    out.append('</cols><sheetData>')
    for ri, row in enumerate(rows, start=1):
        out.append(f'<row r="{ri}">')
        for ci, val in enumerate(row, start=1):
            ref = f'{_col_letter(ci)}{ri}'
            style = ' s="1"' if ri == 1 else ''
            out.append(f'<c r="{ref}" t="inlineStr"{style}><is><t xml:space="preserve">{_esc(val)}</t></is></c>')
        out.append('</row>')
    out.append('</sheetData></worksheet>')
    return ''.join(out)


def write_with_stdlib(path, cols, data, meta):
    meta_rows = [['项', '值']] + [[k, v] for k, v in (meta or [])]
    case_rows = [cols] + data

    content_types = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     '<Default Extension="xml" ContentType="application/xml"/>'
                     '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                     '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                     '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                     '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
                     '</Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            '</Relationships>')
    workbook = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<sheets><sheet name="测试用例" sheetId="1" r:id="rId1"/>'
                '<sheet name="说明" sheetId="2" r:id="rId2"/></sheets></workbook>')
    wb_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
               '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
               '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/>'
               '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
               '</Relationships>')
    styles = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
              '<fonts count="2">'
              '<font><sz val="10"/><name val="Arial"/></font>'
              '<font><b/><sz val="10"/><color rgb="FFFFFFFF"/><name val="Arial"/></font>'
              '</fonts>'
              '<fills count="3"><fill><patternFill patternType="none"/></fill>'
              '<fill><patternFill patternType="gray125"/></fill>'
              '<fill><patternFill patternType="solid"><fgColor rgb="FF305496"/><bgColor indexed="64"/></patternFill></fill></fills>'
              '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
              '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
              '<cellXfs count="2">'
              '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyFont="1"/>'
              '<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1" applyAlignment="1">'
              '<alignment horizontal="center" vertical="center" wrapText="1"/></xf>'
              '</cellXfs>'
              # cellStyles 不能省：缺了它部分解析器会报 "Workbook contains no default style"
              '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
              '</styleSheet>')

    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml', content_types)
        z.writestr('_rels/.rels', rels)
        z.writestr('xl/workbook.xml', workbook)
        z.writestr('xl/_rels/workbook.xml.rels', wb_rels)
        z.writestr('xl/styles.xml', styles)
        z.writestr('xl/worksheets/sheet1.xml', _sheet_xml(case_rows))
        z.writestr('xl/worksheets/sheet2.xml', _sheet_xml(meta_rows))
    return 'stdlib'


# ---------- 主流程 ----------

def main():
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    profile = 'generic'
    if '--profile' in sys.argv:
        i = sys.argv.index('--profile')
        if i + 1 < len(sys.argv):
            profile = sys.argv[i + 1]
    if profile not in ('generic', 'autotest-platform'):
        print(f'未知 profile：{profile}（可选 generic / autotest-platform）')
        return 1
    if not args:
        print(__doc__.strip().splitlines()[-3].strip())
        return 1
    src = Path(args[0])
    if not src.is_file():
        print(f'读不了文件：{src}')
        return 1
    out = Path(sys.argv[sys.argv.index('--out') + 1]) if '--out' in sys.argv else src.with_suffix('.xlsx')

    header, rows = parse_case_table(src.read_text(encoding='utf-8'))
    if not header:
        print(f'在 {src} 里没找到用例表（表头需同时包含「编号」与「判据」）')
        return 1
    cols, data = build_rows(header, rows, profile)

    by_scenario, by_class, verified = {}, {}, 0
    for r in rows:
        sc = cell(r, header, '场景') or '未知'
        by_scenario[sc] = by_scenario.get(sc, 0) + 1
        tc = cell(r, header, '可测性') or '未知'
        by_class[tc] = by_class.get(tc, 0) + 1
        if VERIFIED_RE.search(cell(r, header, '证据/关联')):
            verified += 1
    meta = [
        ('来源文件', str(src)),
        ('生成时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
        ('列结构 profile', profile),
        ('用例总数', len(data)),
        ('按场景', ' · '.join(f'{k} {v}' for k, v in by_scenario.items())),
        ('按可测性', ' · '.join(f'{k} {v}' for k, v in by_class.items())),
        ('已真机验证通过', f'{verified} / {len(data)}'),
        ('生成命令', f'python export_cases_xlsx.py {src.name} --profile {profile}'),
        ('重要提醒', 'AutoTestAgent 的 Excel 导入器只认它那 10 列，判据列会被忽略 —— 导入后用例仍是"无判据"状态'),
    ]

    try:
        engine = write_with_openpyxl(out, cols, data, meta)
    except ImportError:
        engine = write_with_stdlib(out, cols, data, meta)

    print(f'已写出 {len(data)} 条用例 → {out}（引擎：{engine}）')
    print(f'列：{" | ".join(cols)}')
    if engine == 'stdlib':
        print('提示：本机 python 没有 openpyxl，已用纯标准库写入（内容相同、格式较简）。'
              '想要更好的格式：pip install openpyxl，或改用装了解释器（如 py -3.10）。')
    if profile == 'autotest-platform':
        print('提醒：平台的导入器不解析「判据」列 —— 导入后这些用例在平台里仍会被判为"无判据"。')
    return 0


if __name__ == '__main__':
    sys.exit(main())
